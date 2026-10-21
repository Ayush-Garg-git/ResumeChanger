"""
Reads a DOCX resume and lists its paragraphs, in reading order, with IDs.

Why IDs matter
--------------
In Milestone 3 the AI will say things like "paragraph p14 could be shorter".
In Milestone 5 the editor must find that same p14 in the original file.
Both use `iter_paragraphs()` below, which always walks the document in the
same order, so p14 always means the same paragraph.

Every paragraph gets a number, including empty ones. We hide empty ones
from the output, but they still use up a number. That keeps the IDs
identical between reading and editing.

How Word stores text (simplified)
---------------------------------
    document body
      paragraph            <- one line or bullet
        run "Built a "     <- a stretch of text with one formatting
        run "Python"       <- e.g. bold
        run " dashboard"
      table
        row > cell > paragraph > runs ...

What we don't read yet (reported as warnings)
---------------------------------------------
- Text boxes and shapes: some templates put the name or a sidebar there.
- Headers and footers: some templates put contact details there.
"""

import re
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Iterator

from docx import Document
from docx.document import Document as DocumentObject
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

# ---------------------------------------------------------------------------
# Section names
# ---------------------------------------------------------------------------
# Words people use as resume headings, mapped to one standard name.
# This only labels paragraphs to make the outline readable. Nothing is
# changed in the document because of it.
SECTION_KEYWORDS = {
    "Summary": [
        "summary", "professional summary", "profile", "professional profile",
        "objective", "career objective", "about me", "about",
    ],
    "Skills": [
        "skills", "technical skills", "key skills", "core skills",
        "core competencies", "competencies", "tools", "technologies",
        "tools and technologies", "tech stack",
    ],
    "Experience": [
        "experience", "work experience", "professional experience",
        "employment", "employment history", "work history",
        "internship", "internships", "internship experience",
    ],
    "Projects": [
        "projects", "academic projects", "personal projects", "key projects",
        "project experience",
    ],
    "Education": [
        "education", "academic background", "academics", "qualifications",
        "educational qualifications",
    ],
    "Certifications": [
        "certifications", "certificates", "certification", "courses",
        "licenses and certifications", "training",
    ],
    "Achievements": [
        "achievements", "awards", "accomplishments", "honours", "honors",
        "awards and achievements",
    ],
    "Other": [
        "extracurricular", "extracurricular activities", "activities",
        "positions of responsibility", "volunteering", "volunteer experience",
        "leadership", "publications", "languages", "interests", "hobbies",
    ],
}

_KEYWORD_TO_SECTION = {
    kw: section for section, kws in SECTION_KEYWORDS.items() for kw in kws
}

# Paragraphs before the first heading are usually name and contact details.
TOP_SECTION = "Header"


@dataclass
class ResumeParagraph:
    id: str            # "p14": stable reference used by the AI and the editor
    text: str
    section: str       # e.g. "Skills"; a best guess, not a judgment
    is_heading: bool
    is_list_item: bool
    in_table: bool
    style: str         # Word style name, e.g. "List Bullet"


@dataclass
class ResumeOutline:
    paragraphs: list[ResumeParagraph] = field(default_factory=list)
    sections: list[str] = field(default_factory=list)  # in order of appearance
    warnings: list[str] = field(default_factory=list)

    def to_dict(self) -> dict:
        return asdict(self)


# ---------------------------------------------------------------------------
# Walking the document
# ---------------------------------------------------------------------------
def _iter_block_paragraphs(container, parent, in_table: bool) -> Iterator[tuple[Paragraph, bool]]:
    """Yield paragraphs from a body or table cell, in order, entering tables."""
    for child in container.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, parent), in_table

        elif child.tag == qn("w:tbl"):
            table = Table(child, parent)
            # A merged cell appears several times in row.cells; read it once.
            # We keep the cell elements themselves in the set. (Comparing
            # id() numbers is not safe: Python reuses an id once a temporary
            # object is freed, which once made us skip real cells.)
            seen_cells = set()
            for row in table.rows:
                for cell in row.cells:
                    if cell._tc in seen_cells:
                        continue
                    seen_cells.add(cell._tc)
                    yield from _iter_block_paragraphs(cell._tc, cell, in_table=True)

        elif child.tag == qn("w:sdt"):
            # "Content controls": some templates wrap whole sections in them.
            content = child.find(qn("w:sdtContent"))
            if content is not None:
                yield from _iter_block_paragraphs(content, parent, in_table)


def iter_paragraphs(doc: DocumentObject) -> Iterator[tuple[str, Paragraph, bool]]:
    """
    Yield (paragraph_id, paragraph, in_table) for every paragraph in the body.

    This is the single source of truth for paragraph IDs. The editor in
    Milestone 5 will use this same function.
    """
    body = doc.element.body
    for index, (paragraph, in_table) in enumerate(
        _iter_block_paragraphs(body, doc._body, in_table=False)
    ):
        yield f"p{index}", paragraph, in_table


# ---------------------------------------------------------------------------
# Recognising headings and lists
# ---------------------------------------------------------------------------
def _normalise(text: str) -> str:
    text = text.strip().rstrip(":").strip().lower()
    text = text.replace("&", "and")
    return re.sub(r"\s+", " ", text)


def _section_for_heading(paragraph: Paragraph, text: str) -> str | None:
    """Return a section name if this paragraph looks like a heading, else None."""
    key = _normalise(text)

    # 1. Known resume heading words, whatever the formatting.
    if key in _KEYWORD_TO_SECTION:
        return _KEYWORD_TO_SECTION[key]

    # 2. Word's own heading styles, for headings we don't have in our list.
    style = paragraph.style.name if paragraph.style is not None else ""
    if style.startswith("Heading") and len(key.split()) <= 5:
        return text.strip().rstrip(":")

    return None


def _is_list_item(paragraph: Paragraph) -> bool:
    ppr = paragraph._p.pPr
    if ppr is not None and ppr.numPr is not None:
        return True
    style = paragraph.style.name if paragraph.style is not None else ""
    return "List" in style


# ---------------------------------------------------------------------------
# Warnings about content we can't see
# ---------------------------------------------------------------------------
def _collect_warnings(doc: DocumentObject) -> list[str]:
    warnings = []

    body_xml = doc.element.body
    text_boxes = body_xml.findall(".//" + qn("w:txbxContent"))
    if text_boxes:
        warnings.append(
            f"This resume has {len(text_boxes)} text box(es) or shape(s). "
            "Text inside them is not read yet, so it won't get suggestions."
        )

    header_has_text = any(
        p.text.strip()
        for section in doc.sections
        for part in (section.header, section.footer)
        if not part.is_linked_to_previous
        for p in part.paragraphs
    )
    if header_has_text:
        warnings.append(
            "This resume has text in the page header or footer. "
            "That text is not read yet, so it won't get suggestions."
        )

    return warnings


# ---------------------------------------------------------------------------
# Main entry point
# ---------------------------------------------------------------------------
def read_resume(path: Path) -> ResumeOutline:
    doc = Document(str(path))
    outline = ResumeOutline(warnings=_collect_warnings(doc))

    current_section = TOP_SECTION
    for pid, paragraph, in_table in iter_paragraphs(doc):
        text = paragraph.text
        if not text.strip():
            continue  # skipped, but its ID number is still used up

        heading_section = _section_for_heading(paragraph, text)
        if heading_section:
            current_section = heading_section

        if current_section not in outline.sections:
            outline.sections.append(current_section)

        outline.paragraphs.append(
            ResumeParagraph(
                id=pid,
                text=text,
                section=current_section,
                is_heading=heading_section is not None,
                is_list_item=_is_list_item(paragraph),
                in_table=in_table,
                style=paragraph.style.name if paragraph.style is not None else "",
            )
        )

    if not outline.paragraphs:
        outline.warnings.append(
            "No text was found in the main body of this document."
        )

    return outline
