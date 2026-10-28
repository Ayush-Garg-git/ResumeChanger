"""
A factual ATS ("applicant tracking system") check. No AI, no score.

Recruiters' systems read a resume as plain text, then search it for the
job's keywords. This module reports two things we can actually verify:

1. Keywords: which tool-like terms from the job description already appear
   somewhere in the resume, and which don't. Missing ones are NEVER added
   for the user; the page asks about them instead.
2. Formatting: features that some ATS read badly (tables, text boxes,
   columns, contact details only in the page header, unusual headings).

We don't give a percentage "ATS score": every vendor parses differently,
and a number we can't back up would only mislead.
"""

import re
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

from app.docx_reader import SECTION_KEYWORDS, ResumeOutline, TOP_SECTION
from app.suggestion_checks import _PLAIN_WORDS

MAX_KEYWORDS = 30

# Words that are capitalised in job ads for layout reasons (titles, field
# names), not because they are skills.
_NOT_KEYWORDS = set("""
about role job title department type mode location work remote hybrid onsite on-site office
internship intern interns full-time part-time contract permanent engineering engineer engineers
developer developers software hardware responsibilities responsibility required requirements
skills skill eligibility qualifications qualification offer offers application process apply
candidates candidate students student freshers fresher company team teams experience years year
india noida delhi gurugram gurgaon bangalore bengaluru mumbai pune hyderabad chennai pvt ltd
limited private inc llc corp co corporation technologies solutions services we our you your what who why how the and or nice have plus preferred
benefits salary ctc stipend description summary overview details opportunity opportunities mentorship insurance health annual budget
example analyst analysts manager managers executive associate specialist consultant
coordinator officer assistant trainee lead senior junior head director
strong good excellent solid basic advanced proven deep working knowledge familiarity proficiency
proficient exposure ability understanding competitive negotiable immediate joining urgent
monday tuesday wednesday thursday friday saturday sunday january february march april may june
july august september october november december us usa uk uae singapore dubai lpa inr usd lakh lakhs
ahmedabad kolkata jaipur chandigarh indore kochi lucknow bhopal nagpur coimbatore mohali
""".split())

_EMAIL = re.compile(r"[\w.+-]+@[\w-]+(\.[\w-]+)+")
_NUMBER_RUN = re.compile(r"\+?\d[\d\s().-]{8,}\d")

# Version numbers that don't change the skill: CSS3 is CSS.
_VERSIONED = {"css3": "css", "html5": "html", "python3": "python"}


def _has_phone(text: str) -> bool:
    """A run of 10+ digits. Date ranges like "2019 - 2023" have only 8."""
    return any(sum(c.isdigit() for c in m.group()) >= 10 for m in _NUMBER_RUN.finditer(text))


def _forms(word: str) -> set[str]:
    """Ways a keyword may be written in a resume: APIs/API, CSS3/CSS."""
    w = word.casefold().rstrip(".")
    forms = {w, _VERSIONED.get(w, w)}
    if w.endswith("s") and not w.endswith("ss") and len(w) >= 4:
        forms.add(w[:-1])
    return forms


_WORD = re.compile(r"[A-Za-z][A-Za-z0-9+#.-]*[A-Za-z0-9+#]|[A-Za-z]")


def _resume_index(text: str) -> tuple[set[str], set[str]]:
    """(every word in any form, words written with a capital letter)."""
    words, capitalised = set(), set()
    for word in _WORD.findall(text):
        for piece in [word, *word.split("-")]:
            if not piece:
                continue
            words |= _forms(piece)
            if piece[0].isupper():
                capitalised.add(piece.casefold().rstrip("."))
    return words, capitalised


def _present(term: str, words: set[str], capitalised: set[str]) -> bool:
    """
    Is this keyword (or every word of a phrase) in the resume? "React.js"
    counts if the resume says React with a capital, but "Next.js" never
    counts just because the word "next" happens to be there.
    """
    for part in term.split():
        p = part.casefold().rstrip(".")
        if _forms(p) & words:
            continue
        if p.endswith(".js") and p[:-3] in capitalised:
            continue
        return False
    return True


def job_keywords(job_description: str) -> list[str]:
    """
    Tool-like terms from the job, as written there, in order of first
    appearance: acronyms (SQL), mixed case (GitHub, Wi-Fi), digits or
    symbols (C++, STM32, Node.js), or a capitalised word mid-sentence that
    the ad never writes in lower case (React, Altium). Terms written next
    to each other stay together: "Power BI", "REST APIs".
    """
    lowercase_words = {w.casefold() for w in re.findall(r"\b[a-z][a-z0-9+#.-]*\b", job_description)}
    phrases: list[list[str]] = []
    last_end = -2
    for match in _WORD.finditer(job_description):
        word = match.group().rstrip(".")
        term = word.casefold()
        before = job_description[: match.start()].rstrip()
        # Only real sentence starts and bullets: "Tools: Excel" and
        # "(Verilog or VHDL)" still count Excel and Verilog.
        sentence_start = not before or before[-1] in ".!?\n•●○▪■◦‣·*-–—"
        strong = (
            any(c.isdigit() or c in "+#." for c in word)
            or (len(word) >= 2 and word.replace("-", "").isupper())
            or any(c.isupper() for c in word[1:])
        )
        title_case_mid = word[0].isupper() and not sentence_start
        keep = (
            len(term) >= 2
            and (strong or (title_case_mid and term not in lowercase_words))
            and term not in _NOT_KEYWORDS
            and term not in _PLAIN_WORDS
        )
        if not keep:
            continue
        joined = job_description[last_end:match.start()] == " " and phrases
        if joined:
            phrases[-1].append(word)
        else:
            phrases.append([word])
        last_end = match.start() + len(word)
    found, seen = [], set()
    for parts in phrases:
        phrase = " ".join(parts)
        if phrase.casefold() not in seen:
            seen.add(phrase.casefold())
            found.append(phrase)
    return found[:MAX_KEYWORDS]


def keyword_coverage(outline: ResumeOutline, job_description: str) -> dict:
    words, capitalised = _resume_index(" ".join(p.text for p in outline.paragraphs))
    found, missing = [], []
    for term in job_keywords(job_description):
        (found if _present(term, words, capitalised) else missing).append(term)
    return {"found": found, "missing": missing}


# ---------------------------------------------------------------------------
# Formatting
# ---------------------------------------------------------------------------
_STANDARD_HEADINGS = {kw for kws in SECTION_KEYWORDS.values() for kw in kws}


def _heading_key(text: str) -> str:
    """Normalised like the resume reader: "Awards & Achievements" -> "awards and achievements"."""
    return re.sub(r"\s+", " ", text.replace("&", " and ")).strip().casefold().rstrip(":").strip()


def _finding(level: str, text: str) -> dict:
    return {"level": level, "text": text}


def formatting_findings(path: Path, outline: ResumeOutline) -> list[dict]:
    doc = Document(str(path))
    body = doc.element.body
    findings = []

    in_tables = sum(1 for p in outline.paragraphs if p.in_table)
    if in_tables:
        findings.append(_finding(
            "warn",
            f"{in_tables} line(s) are inside a table. Some ATS read tables out of order "
            "or skip them; a single-column layout is safest.",
        ))

    if body.findall(".//" + qn("w:txbxContent")):
        findings.append(_finding(
            "warn", "Some text is inside text boxes or shapes. Many ATS skip it entirely.",
        ))

    columns = [
        c for c in body.iter(qn("w:cols"))
        if (c.get(qn("w:num")) or "1").isdigit() and int(c.get(qn("w:num")) or "1") > 1
    ]
    if columns:
        findings.append(_finding(
            "warn", "Part of the page uses multiple columns. Some ATS mix up the reading order.",
        ))

    header_text = " ".join(
        p.text
        for section in doc.sections
        for part in (
            section.header, section.footer, section.first_page_header,
            section.first_page_footer, section.even_page_header, section.even_page_footer,
        )
        if not part.is_linked_to_previous
        for p in part.paragraphs
    )
    body_text = " ".join(p.text for p in outline.paragraphs)
    if _EMAIL.search(header_text) and not _EMAIL.search(body_text):
        findings.append(_finding(
            "warn", "Your email is only in the page header. Some ATS ignore headers and footers.",
        ))
    elif not _EMAIL.search(body_text) and not _EMAIL.search(header_text):
        findings.append(_finding("warn", "We couldn't find an email address in your resume."))
    if not _has_phone(body_text) and not _has_phone(header_text):
        findings.append(_finding("info", "We couldn't find a phone number in your resume."))

    if body.findall(".//" + qn("w:drawing")):
        findings.append(_finding(
            "info", "Images and icons are ignored by ATS. Make sure no information exists only as a picture.",
        ))

    unusual = [
        p.text.strip() for p in outline.paragraphs
        if p.is_heading and p.section != TOP_SECTION
        and _heading_key(p.text) not in _STANDARD_HEADINGS
    ]
    for heading in unusual[:3]:
        findings.append(_finding(
            "info",
            f"The heading '{heading}' isn't a standard one. ATS sort content by headings like "
            "Summary, Experience, Projects, Skills, Education; a standard name may parse better.",
        ))

    if not any(f["level"] == "warn" for f in findings):
        findings.insert(0, _finding(
            "ok", "No layout problems found: the text is in the main body, without text boxes or columns.",
        ))
    return findings


def ats_report(path: Path, outline: ResumeOutline, job_description: str) -> dict:
    return {
        "keywords": keyword_coverage(outline, job_description),
        "formatting": formatting_findings(path, outline),
    }
