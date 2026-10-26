"""
Tests for the preview and the page limit (Milestone 6).

Rendering tests need LibreOffice. They run where it's installed (the Docker
image, or a machine that has it) and are skipped elsewhere, e.g. a Windows
laptop without LibreOffice.
"""

import json

import pytest
from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from fastapi.testclient import TestClient

from tests.helpers import ai_post

from app import analyzer, renderer, storage
from app.config import settings
from app.docx_editor import Change, apply_changes, mark_original
from app.docx_reader import iter_paragraphs
from app.main import app
from app.page_fit import page_fit

client = TestClient(app)
needs_libreoffice = pytest.mark.skipif(not renderer.is_available(), reason="LibreOffice not installed")

P10 = "Built a Power BI dashboard used by regional managers to track weekly sales."
P11 = "Automated a manual Excel report with Python, cutting preparation time from a full day to under an hour."
P13 = "Inventory Forecast: Forecasted monthly stock using past sales in Python."
P14 = "Portfolio Website: Personal site built with HTML and CSS."


# --- Page limit options (no rendering needed) -------------------------------

ANALYSIS = {"suggestions": [
    {"id": "s1", "paragraph_id": "p10", "kind": "reword", "original_text": P10, "proposed_text": "x", "reason": "r"},
    {"id": "s2", "paragraph_id": "p11", "kind": "shorten", "original_text": P11, "proposed_text": "short", "reason": "r"},
    {"id": "s3", "paragraph_id": "p13", "kind": "review", "original_text": P13, "proposed_text": None, "reason": "r"},
    {"id": "s4", "paragraph_id": "p14", "kind": "consider_removing", "original_text": P14, "proposed_text": None, "reason": "r"},
]}


def test_fits_means_no_options():
    result = page_fit(ANALYSIS, {}, pages=1, max_pages=1)
    assert result["fits"] and result["options"] == []


def test_no_limit_always_fits():
    assert page_fit(ANALYSIS, {}, pages=4, max_pages=None)["fits"]


def test_over_the_limit_lists_options_in_order():
    result = page_fit(ANALYSIS, {}, pages=2, max_pages=1)
    assert not result["fits"] and result["over_by"] == 1
    assert [o["type"] for o in result["options"]] == [
        "remove", "shorten", "review", "adjust_layout", "keep_longer",
    ]
    # Rewording doesn't save space, so it's not offered.
    assert all(o.get("suggestion_id") != "s1" for o in result["options"])


def test_already_applied_options_are_not_offered_again():
    decisions = {"s4": {"effect": "remove"}, "s2": {"effect": "change"}}
    types = [o["type"] for o in page_fit(ANALYSIS, decisions, 2, 1)["options"]]
    assert "remove" not in types and "shorten" not in types


def test_lines_the_user_kept_are_offered_but_marked():
    option = page_fit(ANALYSIS, {"s4": {"effect": "keep"}}, 2, 1)["options"][0]
    assert option["type"] == "remove" and option["previously_kept"] is True


def test_nothing_is_ever_removed_automatically():
    # page_fit only describes options. It returns data; it changes nothing.
    decisions = {}
    page_fit(ANALYSIS, decisions, 3, 1)
    assert decisions == {}


# --- Highlighted copies for the preview --------------------------------------

def test_after_copy_highlights_only_changed_words(sample_resume, tmp_path):
    out = tmp_path / "after.docx"
    apply_changes(sample_resume, out, [Change("s1", "p6", "change", "Languages: Python, SQL", "Languages: SQL, Python")], highlight=True)
    runs = {i: p for i, p, _ in iter_paragraphs(Document(str(out)))}["p6"].runs
    assert runs[0].font.highlight_color is None                # bold label untouched
    assert runs[1].font.highlight_color == WD_COLOR_INDEX.YELLOW


def test_download_copy_has_no_highlight(sample_resume, tmp_path):
    out = tmp_path / "clean.docx"
    apply_changes(sample_resume, out, [Change("s1", "p6", "change", "Languages: Python, SQL", "Languages: SQL, Python")])
    para = {i: p for i, p, _ in iter_paragraphs(Document(str(out)))}["p6"]
    assert all(r.font.highlight_color is None for r in para.runs)


def test_before_copy_marks_changed_and_removed_lines(sample_resume, tmp_path):
    out = tmp_path / "before.docx"
    mark_original(sample_resume, out, [{"paragraph_id": "p10", "how": "changed"}, {"paragraph_id": "p14", "how": "removed"}])
    paras = {i: p for i, p, _ in iter_paragraphs(Document(str(out)))}
    assert paras["p10"].runs[0].font.highlight_color == WD_COLOR_INDEX.YELLOW
    assert paras["p14"].runs[0].font.highlight_color == WD_COLOR_INDEX.RED
    assert paras["p14"].text == P14                            # text not changed
    assert paras["p9"].runs[0].font.highlight_color is None    # others untouched


# --- Rendering (real LibreOffice) -------------------------------------------

@needs_libreoffice
def test_render_counts_pages_and_draws_images(sample_resume, tmp_path):
    result = renderer.render(sample_resume, tmp_path / "out", "before")
    assert result.pages == 1
    assert [p.name for p in result.images] == ["before-1.png"]
    assert result.images[0].read_bytes()[:8] == b"\x89PNG\r\n\x1a\n"


@needs_libreoffice
def test_render_counts_two_pages(tmp_path):
    doc = Document()
    for i in range(90):
        doc.add_paragraph(f"Line {i}: built and maintained internal reporting tools for teams.")
    path = tmp_path / "long.docx"
    doc.save(str(path))
    assert renderer.render(path, tmp_path / "out", "after").pages >= 2


def test_render_reports_missing_libreoffice(sample_resume, tmp_path, monkeypatch):
    monkeypatch.setattr(renderer.shutil, "which", lambda name: None)
    with pytest.raises(renderer.RenderError, match="not installed"):
        renderer.render(sample_resume, tmp_path, "before")


# --- Routes -----------------------------------------------------------------

def _answer():
    s = lambda pid, kind, orig, prop: {"paragraph_id": pid, "kind": kind, "original_text": orig,  # noqa: E731
                                       "proposed_text": prop, "reason": "r", "requirement": None}
    return {"role_title": "Data Analyst", "fit_summary": "ok", "requirements": [], "questions_for_user": [],
            "no_changes_needed": False, "suggestions": [
                s("p10", "reword", P10, "Built a Power BI dashboard that regional managers used to track weekly sales."),
                s("p11", "shorten", P11, "Automated a manual Excel report with Python, cutting preparation from a day to under an hour."),
                s("p14", "consider_removing", P14, None),
            ]}


@pytest.fixture
def applied_session(sample_resume, monkeypatch):
    monkeypatch.setattr(analyzer, "call_model", lambda prompt, schema: json.dumps(_answer()))
    monkeypatch.setattr(settings, "gemini_api_key", "test-key")
    files = {"file": ("r.docx", sample_resume.read_bytes(), "application/octet-stream")}
    sid = client.post("/api/resume", files=files).json()["session_id"]
    job = "Data Analyst role: Power BI dashboards, SQL, weekly sales reporting for regional managers."
    ai_post(client, f"/api/resume/{sid}/analyze", json={"job_description": job, "consent": True})
    aid = storage.load_json(sid, "analysis.json")["analysis_id"]
    client.put(f"/api/resume/{sid}/suggestions/s1", json={"action": "accept", "analysis_id": aid})
    assert client.post(f"/api/resume/{sid}/apply").json()["download_ready"]
    return sid


def test_preview_needs_a_refined_file(sample_resume):
    files = {"file": ("r.docx", sample_resume.read_bytes(), "application/octet-stream")}
    sid = client.post("/api/resume", files=files).json()["session_id"]
    assert client.get(f"/api/resume/{sid}/preview").status_code == 404


def test_preview_without_libreoffice_says_so(applied_session, monkeypatch):
    monkeypatch.setattr(renderer, "is_available", lambda: False)
    body = client.get(f"/api/resume/{applied_session}/preview").json()
    assert body["available"] is False


@needs_libreoffice
def test_preview_page_fit_and_cleanup(applied_session):
    sid = applied_session
    body = client.get(f"/api/resume/{sid}/preview").json()
    assert body["available"] and body["before"]["pages"] == 1 and body["after"]["pages"] == 1

    image = client.get(body["after"]["images"][0])
    assert image.status_code == 200 and image.headers["content-type"] == "image/png"

    # Marked copies are never left behind; only images are kept.
    assert not list(storage.preview_dir(sid).glob("*.docx"))

    fit = client.get(f"/api/resume/{sid}/page-fit", params={"max_pages": 1}).json()
    assert fit["fits"] is True

    # Changing a decision withdraws the preview with the file.
    aid = storage.load_json(sid, "analysis.json")["analysis_id"]
    client.put(f"/api/resume/{sid}/suggestions/s3", json={"action": "accept", "analysis_id": aid})
    assert not storage.preview_dir(sid).exists()


def test_preview_image_names_are_checked(applied_session):
    assert client.get(f"/api/resume/{applied_session}/preview/..%2Foriginal.docx").status_code in (400, 404)
    assert client.get(f"/api/resume/{applied_session}/preview/original.docx").status_code == 400


def test_page_fit_limits(applied_session):
    assert client.get(f"/api/resume/{applied_session}/page-fit", params={"max_pages": 0}).status_code == 400
