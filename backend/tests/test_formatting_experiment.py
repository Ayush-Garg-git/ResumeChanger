"""
The formatting experiment (preparing for Milestone 5).

Question: if we change a few words in a paragraph, does the formatting
survive? This file records the answer as tests.

Setup: the line "Languages: Python, SQL" has two runs:
    run 0: "Languages: "   (bold)
    run 1: "Python, SQL"   (normal)

Finding 1: setting paragraph.text replaces ALL runs with one plain run,
           so the bold label is lost. We must never edit that way.
Finding 2: changing the text of only the run that contains the words
           keeps every other run exactly as it was.
"""

from docx import Document

from app.docx_reader import iter_paragraphs


def _languages_paragraph(doc):
    return next(p for _, p, _ in iter_paragraphs(doc) if p.text.startswith("Languages:"))


def test_naive_edit_loses_bold(sample_resume, tmp_path):
    doc = Document(str(sample_resume))
    para = _languages_paragraph(doc)
    assert para.runs[0].bold is True  # starts bold

    para.text = "Languages: Python, SQL, R"  # the naive way

    out = tmp_path / "naive.docx"
    doc.save(str(out))
    para_after = _languages_paragraph(Document(str(out)))

    assert len(para_after.runs) == 1
    assert para_after.runs[0].bold is not True  # bold is gone


def test_run_level_edit_keeps_bold(sample_resume, tmp_path):
    doc = Document(str(sample_resume))
    para = _languages_paragraph(doc)

    # Change only the run that holds the words we're editing.
    target = next(r for r in para.runs if "Python, SQL" in r.text)
    target.text = target.text.replace("Python, SQL", "Python, SQL, R")

    out = tmp_path / "careful.docx"
    doc.save(str(out))
    para_after = _languages_paragraph(Document(str(out)))

    assert para_after.text == "Languages: Python, SQL, R"
    assert para_after.runs[0].text == "Languages: "
    assert para_after.runs[0].bold is True  # bold survived
    assert para_after.runs[1].bold is not True
