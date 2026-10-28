"""
Tests for whole-resume tailoring: section reviews, the new "tailor" and
"remove_items" kinds, the stricter honesty checks that come with them,
and the ATS check.
"""

import io
import json

import pytest
from docx import Document
from fastapi.testclient import TestClient

from tests.helpers import ai_post

from app import analyzer
from app.analyzer import paragraphs_to_send, require_all
from app.ats import ats_report, job_keywords
from app.config import settings
from app.docx_reader import read_resume
from app.main import app
from app.schemas import ResumeAnalysis, SectionReview, Suggestion
from app.suggestion_checks import (
    check_section_reviews,
    check_suggestions,
    invented_terms,
    invented_words,
    removal_problem,
    word_form,
)
from app.suggestion_checks import _tokens

client = TestClient(app)

JOB = (
    "Data Analyst at Example Corp. You will build Power BI dashboards, write SQL, "
    "and present weekly sales insights to regional managers. Tableau is a plus."
)
P4 = "Analyst with two years of experience turning sales and operations data into weekly reports for business teams."
P10 = "Built a Power BI dashboard used by regional managers to track weekly sales."
P11 = "Automated a manual Excel report with Python, cutting preparation time from a full day to under an hour."


def _s(pid, original, proposed, kind):
    return Suggestion(paragraph_id=pid, kind=kind, original_text=original,
                      proposed_text=proposed, reason="Relevant.", requirement=None)


def _analysis(suggestions=(), reviews=()):
    return ResumeAnalysis(role_title="Data Analyst", fit_summary="Fits.", requirements=[],
                          suggestions=list(suggestions), questions_for_user=[],
                          no_changes_needed=False, section_reviews=list(reviews))


def _check(sample_resume, *suggestions):
    sent = paragraphs_to_send(read_resume(sample_resume))
    return check_suggestions(_analysis(suggestions), sent, JOB)


# --------------------------------------------------------------------------
# Tailoring
# --------------------------------------------------------------------------
def test_tailored_summary_using_resume_facts_is_kept(sample_resume):
    kept, dropped = _check(sample_resume, _s(
        "p4", P4,
        "Analyst with two years of experience building Power BI dashboards and Python reports "
        "on sales data for business teams.", "tailor"))
    assert len(kept) == 1 and kept[0].kind == "tailor", dropped


def test_plural_of_a_resume_word_is_not_treated_as_invented(sample_resume):
    # "dashboards" is in the job; the resume says "dashboard". Same thing.
    kept, dropped = _check(sample_resume, _s(
        "p10", P10, "Built Power BI dashboards that regional managers used to track weekly sales.", "tailor"))
    assert len(kept) == 1, dropped


def test_tailoring_still_cannot_add_a_job_skill_the_resume_lacks(sample_resume):
    kept, dropped = _check(sample_resume, _s(
        "p10", P10, "Built a Power BI and Tableau dashboard used by regional managers to track weekly sales.", "tailor"))
    assert not kept and "tableau" in dropped[0].why


def test_tailoring_cannot_drop_a_result(sample_resume):
    kept, dropped = _check(sample_resume, _s(
        "p9", "Data Analyst, Example Retail Pvt Ltd (2023 to present)",
        "Data Analyst, Example Retail Pvt Ltd", "tailor"))
    assert not kept and "dropped a result" in dropped[0].why


def test_shortening_may_drop_a_number(sample_resume):
    kept, dropped = _check(sample_resume, _s(
        "p9", "Data Analyst, Example Retail Pvt Ltd (2023 to present)",
        "Data Analyst, Example Retail Pvt Ltd", "shorten"))
    assert len(kept) == 1, dropped


def test_tailoring_cannot_inflate(sample_resume):
    kept, dropped = _check(sample_resume, _s(
        "p10", P10, "Led a Power BI dashboard used by regional managers to track weekly sales.", "tailor"))
    assert not kept and "led" in dropped[0].why


# --------------------------------------------------------------------------
# Remove items
# --------------------------------------------------------------------------
def test_removing_a_skill_from_a_list_is_kept(sample_resume):
    kept, dropped = _check(sample_resume, _s("p7", "Tools: Excel, Power BI, Git", "Tools: Power BI, Git", "remove_items"))
    assert len(kept) == 1 and kept[0].kind == "remove_items", dropped


@pytest.mark.parametrize("proposed,why", [
    ("Tools: Power BI, Git, Tableau", "changed or added"),     # sneaks one in
    ("Tools: Git, Power BI", "changed or added"),             # reorders
    ("Software: Power BI, Git", "start of the line"),         # changes the label
    ("Tools: Excel, Power BI, Git.", "removed nothing"),       # removes nothing
])
def test_remove_items_must_only_remove(sample_resume, proposed, why):
    kept, dropped = _check(sample_resume, _s("p7", "Tools: Excel, Power BI, Git", proposed, "remove_items"))
    assert not kept and why in dropped[0].why


def test_remove_items_on_a_sentence_is_refused():
    assert "isn't a simple list" in removal_problem(P10, "Built a Power BI dashboard.")


@pytest.mark.parametrize("original,proposed", [
    ("Proficient: Python, SQL | Beginner: Java, C", "Proficient: Python, SQL | C"),
    ("Wrote unit tests in Python, under supervision of a senior engineer, for the billing team",
     "Wrote unit tests in Python, for the billing team"),
    ("Certifications: AWS, SQL (2023), Python", "Certifications: AWS, Python"),
])
def test_remove_items_cannot_change_meaning(original, proposed):
    assert removal_problem(original, proposed) is not None


def test_remove_items_only_in_the_skills_section(sample_resume):
    kept, dropped = _check(sample_resume, _s(
        "p16", "B.Com, Example University", "B.Com", "remove_items"))
    assert not kept and "skills lists" in dropped[0].why


def test_items_in_brackets_stay_whole():
    original = "Systems: Linux (CLI, configuration, troubleshooting), Windows, networking"
    assert removal_problem(original, "Systems: Linux (CLI, configuration, troubleshooting), networking") is None
    # Editing inside a bracketed item is a change, not a removal of items.
    assert removal_problem(original, "Systems: Linux (CLI, troubleshooting), Windows, networking") is not None


# --------------------------------------------------------------------------
# Word forms: flexible enough for honest tailoring, never a loophole
# --------------------------------------------------------------------------
@pytest.mark.parametrize("a,b", [("dashboards", "dashboard"), ("APIs", "API"), ("debugging", "debugged"),
                                 ("built", "building"), ("wrote", "writing"), ("processes", "process")])
def test_forms_of_the_same_word_match(a, b):
    assert word_form(a) == word_form(b)


@pytest.mark.parametrize("tool,resume_word", [
    ("Excel", "excelled"), ("React", "reacted"), ("Spark", "sparked"), ("Bootstrap", "bootstrapped"),
    ("marketing", "market"), ("programming", "program"), ("accounting", "account"), ("hosting", "host"),
    ("presented", "Present"),  # "Jan 2023 - Present" must not vouch for presenting
])
def test_unrelated_words_with_a_shared_root_do_not_vouch(tool, resume_word):
    assert word_form(tool) != word_form(resume_word)


def test_excelled_in_a_resume_does_not_let_excel_in():
    resume = _tokens("Excelled in mathematics; reacted quickly to incidents; sparked interest in clubs")
    jd = "We need Excel, React and Spark."
    assert {"excel", "react", "spark"} <= set(invented_terms("Built reports in Excel using React and Spark", resume, jd))


def test_a_date_present_does_not_let_presented_in():
    resume = _tokens("Analyst, Example Ltd (2023 - Present). Built a dashboard.")
    assert "presented" in invented_words("Built and presented a dashboard", resume, "Present insights to managers.")


def test_shortening_cannot_drop_shared_credit(sample_resume):
    sent = paragraphs_to_send(read_resume(sample_resume))
    line = sent[[p.id for p in sent].index("p10")]
    from app.suggestion_checks import dropped_shared_credit
    assert dropped_shared_credit("Assisted senior engineers in designing the PCB.", "Designed the PCB.") == ["assisted"]
    assert dropped_shared_credit("Helped the team with testing the module.", "Tested the module.") == ["helped"]
    assert dropped_shared_credit(line.text, line.text) == []


def test_inflation_is_caught_in_any_form(sample_resume):
    for verb in ["Leads", "Managing", "Owns", "Oversees"]:
        kept, dropped = _check(sample_resume, _s(
            "p10", P10, f"{verb} a Power BI dashboard used by regional managers to track weekly sales.", "tailor"))
        assert not kept, verb


# --------------------------------------------------------------------------
# Section reviews
# --------------------------------------------------------------------------
def _review(section, tips=(), relevance="high"):
    return SectionReview(section=section, relevance=relevance, verdict="Fine.", tips=list(tips))


def test_section_reviews_follow_resume_order_and_real_sections(sample_resume):
    sent = paragraphs_to_send(read_resume(sample_resume))
    reviews = check_section_reviews(_analysis(reviews=[
        _review("projects"), _review("Summary"), _review("Hobbies"), _review("Summary"),
    ]), sent, JOB)
    # Case-insensitive match, real names, resume order, no duplicates or made-up sections.
    assert [r["section"] for r in reviews] == ["Summary", "Projects"]


def test_tips_that_suggest_adding_missing_skills_are_dropped(sample_resume):
    sent = paragraphs_to_send(read_resume(sample_resume))
    reviews = check_section_reviews(_analysis(reviews=[_review("Skills", tips=[
        "Add Tableau to your tools.",                      # job skill not in the resume: dropped
        "List Docker under Tools too.",                    # tool-like, in neither: dropped
        "Describe the insights you presented.",            # job claim the resume doesn't make: dropped
        "Move this section above Work Experience.",        # advice only: kept
        "Put Power BI first, since the job leads with it.",
    ])]), sent, JOB)
    assert reviews[0]["tips"] == [
        "Move this section above Work Experience.",
        "Put Power BI first, since the job leads with it.",
    ]


def test_schema_sent_to_the_ai_requires_every_section_review_field():
    schema = require_all(ResumeAnalysis.model_json_schema())
    assert "section_reviews" in schema["required"]
    review = schema["$defs"]["SectionReview"]
    assert set(review["required"]) == {"section", "relevance", "verdict", "tips"}


# --------------------------------------------------------------------------
# ATS check
# --------------------------------------------------------------------------
def test_job_keywords_are_tool_like_terms_as_written():
    jd = ("Job Title: Software Developer Intern. Work Mode: Remote / Hybrid. Familiarity with React.js "
          "and Node.js, HTML, CSS, MySQL or MongoDB, Git and GitHub, Wi-Fi modules, STM32 boards. "
          "You will work with a developer team.")
    terms = job_keywords(jd)
    for expected in ["React.js", "Node.js", "HTML", "CSS", "MySQL", "MongoDB", "GitHub", "Wi-Fi", "STM32"]:
        assert expected in terms
    for noise in ["Developer", "Intern", "Remote", "Hybrid", "Software"]:
        assert noise not in terms


def test_keyword_coverage_matches_versions_and_suffixes(sample_resume, tmp_path):
    outline = read_resume(sample_resume)
    report = ats_report(sample_resume, outline, "We use Power BI, SQL, Python3 dashboards, CSS3 and Tableau.")
    assert "SQL" in report["keywords"]["found"]
    assert "Python3" in report["keywords"]["found"]   # resume says "Python"
    assert "CSS3" in report["keywords"]["found"]      # resume says "CSS"
    assert "Tableau" in report["keywords"]["missing"]


def test_clean_single_column_resume_has_no_layout_warnings(tmp_path):
    doc = Document()
    for text in ["Asha Rao", "asha@example.com | +91 98765 43210", "Experience", "Built reports in Python."]:
        doc.add_paragraph(text)
    path = tmp_path / "clean.docx"
    doc.save(path)
    findings = ats_report(path, read_resume(path), JOB)["formatting"]
    assert findings[0]["level"] == "ok"
    assert not any(f["level"] == "warn" for f in findings)


def test_sample_resume_table_is_reported(sample_resume):
    # The sample keeps education dates in a table: worth knowing for ATS.
    findings = ats_report(sample_resume, read_resume(sample_resume), JOB)["formatting"]
    assert any(f["level"] == "warn" and "inside a table" in f["text"] for f in findings)


def test_two_column_layout_is_reported(two_column_resume):
    findings = ats_report(two_column_resume, read_resume(two_column_resume), JOB)["formatting"]
    assert any("multiple columns" in f["text"] for f in findings)


def test_tables_and_header_only_email_are_flagged(tmp_path):
    doc = Document()
    doc.sections[0].header.paragraphs[0].text = "asha@example.com | +91 98765 43210"
    doc.add_paragraph("Asha Rao")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Skills"
    table.cell(0, 1).text = "Python, SQL"
    doc.add_paragraph("Experience")
    doc.add_paragraph("Built reports in Python.")
    path = tmp_path / "t.docx"
    doc.save(path)
    texts = " ".join(f["text"] for f in ats_report(path, read_resume(path), JOB)["formatting"])
    assert "inside a table" in texts
    assert "only in the page header" in texts


# --------------------------------------------------------------------------
# End to end through the API
# --------------------------------------------------------------------------
WHOLE_RESUME_ANSWER = {
    "role_title": "Data Analyst",
    "fit_summary": "Good fit.",
    "requirements": [],
    "section_reviews": [
        {"section": "Summary", "relevance": "high", "verdict": "Leads with reporting.", "tips": []},
        {"section": "Skills", "relevance": "high", "verdict": "Relevant tools.", "tips": ["Move Skills below Work Experience."]},
        {"section": "Experience", "relevance": "high", "verdict": "Strong match.", "tips": []},
        {"section": "Projects", "relevance": "medium", "verdict": "One relevant project.", "tips": []},
        {"section": "Education", "relevance": "low", "verdict": "Not central.", "tips": []},
        {"section": "Certifications", "relevance": "high", "verdict": "Relevant certificate.", "tips": []},
    ],
    "suggestions": [
        {"paragraph_id": "p4", "kind": "tailor", "original_text": "Analyst with two years of",
         "proposed_text": "Analyst with two years of experience building Power BI dashboards and Python "
                          "reports on sales data for business teams.",
         "reason": "Leads with the job's tools.", "requirement": None},
        {"paragraph_id": "p7", "kind": "remove_items", "original_text": "Tools: Excel, Power BI, Git",
         "proposed_text": "Tools: Power BI, Git", "reason": "Excel adds little here.", "requirement": None},
    ],
    "questions_for_user": ["Have you used Tableau?"],
    "no_changes_needed": False,
}


@pytest.fixture
def whole_resume_ai(monkeypatch):
    monkeypatch.setattr(settings, "gemini_api_key", "test-key")
    monkeypatch.setattr(analyzer, "call_model", lambda prompt, schema, model=None: json.dumps(WHOLE_RESUME_ANSWER))


def _upload(sample_resume):
    files = {"file": ("resume.docx", sample_resume.read_bytes(), "application/octet-stream")}
    return client.post("/api/resume", files=files).json()["session_id"]


def test_analysis_covers_every_section_with_ats(sample_resume, whole_resume_ai):
    sid = _upload(sample_resume)
    result = ai_post(client, f"/api/resume/{sid}/analyze", json={"job_description": JOB, "consent": True}).json()
    assert [s["name"] for s in result["sections"]] == [
        "Summary", "Skills", "Experience", "Projects", "Education", "Certifications"]
    assert [r["section"] for r in result["section_reviews"]] == [s["name"] for s in result["sections"]]
    assert {s["kind"] for s in result["suggestions"]} == {"tailor", "remove_items"}
    assert {"Power BI", "SQL"} <= set(result["ats"]["keywords"]["found"])  # kept together
    assert "Corp" not in result["ats"]["keywords"]["missing"]  # company names aren't skills
    assert result["ats"]["formatting"]


def test_accepted_tailor_and_removal_reach_the_word_file(sample_resume, whole_resume_ai):
    sid = _upload(sample_resume)
    result = ai_post(client, f"/api/resume/{sid}/analyze", json={"job_description": JOB, "consent": True}).json()
    for s in result["suggestions"]:
        r = client.put(f"/api/resume/{sid}/suggestions/{s['id']}",
                       json={"action": "accept", "analysis_id": result["analysis_id"]})
        assert r.status_code == 200, r.json()
    applied = client.post(f"/api/resume/{sid}/apply")
    assert applied.status_code == 200, applied.json()
    data = client.get(f"/api/resume/{sid}/download").content
    texts = [p.text for p in Document(io.BytesIO(data)).paragraphs]
    assert "Tools: Power BI, Git" in texts
    assert any(t.startswith("Analyst with two years of experience building Power BI dashboards") for t in texts)
    assert "Languages: Python, SQL" in texts  # untouched lines stay as they were


def test_adjacent_keywords_stay_together_and_role_words_are_skipped():
    terms = job_keywords("Data Analyst, Example Corp. Build Power BI dashboards with REST APIs in Altium Designer.")
    assert terms[:3] == ["Power BI", "REST APIs", "Altium Designer"]
    assert "Analyst" not in terms


# --- Review follow-ups: ATS keyword quality and layout ---------------------

def test_keywords_after_colons_and_brackets_are_found_and_noise_is_not():
    jd = ("Location: Ahmedabad. Salary: 6 LPA. Tools: Excel, Tableau (Verilog or VHDL).\n"
          "● Develop REST APIs\n● Collaborate with US teams on Monday.\nStrong SQL skills.")
    terms = job_keywords(jd)
    for expected in ["Excel", "Tableau", "Verilog", "VHDL", "REST APIs", "SQL"]:
        assert expected in terms
    for noise in ["Ahmedabad", "LPA", "Develop", "Collaborate", "US", "Monday", "Strong SQL", "Strong"]:
        assert noise not in terms


def _resume_with(tmp_path, *lines):
    doc = Document()
    for line in lines:
        doc.add_paragraph(line)
    path = tmp_path / "r.docx"
    doc.save(path)
    return path


def test_keyword_matching_is_not_fooled_by_common_words(tmp_path):
    path = _resume_with(tmp_path, "Asha Rao", "asha@example.com", "Experience",
                        "Planned the next release of our web portal using REST API calls and React.")
    kw = ats_report(path, read_resume(path), "Familiarity with Next.js, Web3, React.js and REST APIs.")["keywords"]
    assert "Next.js" in kw["missing"] and "Web3" in kw["missing"]
    assert "React.js" in kw["found"] and "REST APIs" in kw["found"]


def test_date_ranges_are_not_mistaken_for_a_phone_number(tmp_path):
    path = _resume_with(tmp_path, "Asha Rao", "asha@example.com", "Experience", "Analyst, 2019 - 2023")
    texts = " ".join(f["text"] for f in ats_report(path, read_resume(path), JOB)["formatting"])
    assert "couldn't find a phone number" in texts


def test_headings_with_ampersand_count_as_standard(tmp_path):
    path = _resume_with(tmp_path, "Asha Rao", "asha@example.com | +91 98765 43210",
                        "Awards & Achievements", "Best Analyst award.")
    texts = " ".join(f["text"] for f in ats_report(path, read_resume(path), JOB)["formatting"])
    assert "isn't a standard one" not in texts


def test_shortening_that_changes_who_did_the_work_is_dropped(tmp_path):
    path = _resume_with(tmp_path, "Asha Rao", "asha@example.com", "Experience",
                        "Assisted senior engineers in designing the billing dashboard.")
    sent = paragraphs_to_send(read_resume(path))
    line = next(p for p in sent if p.text.startswith("Assisted"))
    kept, dropped = check_suggestions(_analysis([_s(
        line.id, line.text, "Designed the billing dashboard.", "shorten")]), sent, "Design dashboards.")
    assert not kept and "assisted" in dropped[0].why
