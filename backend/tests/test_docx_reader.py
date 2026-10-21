"""Tests for reading a resume into paragraphs and sections."""

from docx import Document

from app.docx_reader import iter_paragraphs, read_resume


def _by_text(outline, start):
    return next(p for p in outline.paragraphs if p.text.startswith(start))


def test_sections_are_found_in_order(sample_resume):
    outline = read_resume(sample_resume)

    assert outline.sections == [
        "Header", "Summary", "Skills", "Experience",
        "Projects", "Education", "Certifications",
    ]


def test_both_heading_styles_are_recognised(sample_resume):
    outline = read_resume(sample_resume)

    # Word "Heading 1" style
    assert _by_text(outline, "Work Experience").is_heading
    # Plain bold text in capitals
    assert _by_text(outline, "SKILLS").is_heading
    # Plain bold text with a trailing colon
    assert _by_text(outline, "Certifications:").is_heading


def test_a_skills_line_is_content_not_a_heading(sample_resume):
    # "Languages: Python, SQL" starts like a label but is content.
    outline = read_resume(sample_resume)
    line = _by_text(outline, "Languages:")

    assert not line.is_heading
    assert line.section == "Skills"


def test_bullets_and_tables_are_marked(sample_resume):
    outline = read_resume(sample_resume)

    assert _by_text(outline, "Built a Power BI").is_list_item
    edu = _by_text(outline, "B.Com")
    assert edu.in_table
    assert edu.section == "Education"


def test_every_table_cell_is_read_once(sample_resume):
    outline = read_resume(sample_resume)
    table_texts = [p.text for p in outline.paragraphs if p.in_table]

    assert table_texts == [
        "B.Com, Example University", "2019 to 2022",
        "Class XII, Example School", "2019",
    ]


def test_merged_cells_are_not_repeated(tmp_path):
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    merged = table.cell(0, 0).merge(table.cell(0, 1))
    merged.text = "Merged heading cell"
    table.cell(1, 0).text = "Left"
    table.cell(1, 1).text = "Right"
    path = tmp_path / "merged.docx"
    doc.save(str(path))

    texts = [p.text for p in read_resume(path).paragraphs]

    assert texts == ["Merged heading cell", "Left", "Right"]


def test_nothing_is_lost(sample_resume):
    # Every non-empty paragraph python-docx can see must be in our outline.
    doc = Document(str(sample_resume))
    all_text = [para.text for _, para, _ in iter_paragraphs(doc) if para.text.strip()]

    assert [p.text for p in read_resume(sample_resume).paragraphs] == all_text
    assert "Class XII, Example School" in all_text


def test_empty_paragraphs_are_hidden_but_keep_their_number(sample_resume):
    outline = read_resume(sample_resume)
    ids = [p.id for p in outline.paragraphs]

    assert all(p.text.strip() for p in outline.paragraphs)
    # p0 name, p1 contact, p2 is the empty line, so the summary heading is p3.
    assert ids[:3] == ["p0", "p1", "p3"]


def test_ids_point_to_the_same_paragraph_when_read_again(sample_resume):
    # The editor (Milestone 5) relies on this: walking the document again
    # must give the same paragraph for the same ID.
    outline = read_resume(sample_resume)
    doc = Document(str(sample_resume))
    by_id = {pid: para.text for pid, para, _ in iter_paragraphs(doc)}

    for p in outline.paragraphs:
        assert by_id[p.id] == p.text


def test_reading_does_not_change_the_file(sample_resume):
    before = sample_resume.read_bytes()
    read_resume(sample_resume)
    assert sample_resume.read_bytes() == before


def test_header_text_produces_a_warning(sample_resume):
    doc = Document(str(sample_resume))
    doc.sections[0].header.paragraphs[0].text = "Riya Mehta | riya.mehta@example.com"
    doc.save(str(sample_resume))

    outline = read_resume(sample_resume)

    assert any("header or footer" in w for w in outline.warnings)


# --- Real-world layouts -----------------------------------------------------
# Found while testing a real resume converted from PDF: the name was styled
# as Heading 1 and the SUMMARY label was stored before it in the file.

def test_name_in_heading_style_is_not_a_section(two_column_resume):
    outline = read_resume(two_column_resume)
    name = _by_text(outline, "ARJUN RAO")

    assert not name.is_heading
    assert name.section == "Header"
    assert "ARJUN RAO" not in outline.sections


def test_header_block_stays_together(two_column_resume):
    outline = read_resume(two_column_resume)

    assert _by_text(outline, "Computer Science Undergraduate").section == "Header"
    assert _by_text(outline, "Bengaluru").section == "Header"


def test_heading_stored_before_the_name_waits_for_its_content(two_column_resume):
    outline = read_resume(two_column_resume)

    assert _by_text(outline, "SUMMARY").section == "Summary"
    assert _by_text(outline, "Developer who builds").section == "Summary"
    # Order as a reader sees it on the page: header first.
    assert outline.sections == ["Header", "Summary", "Education", "COMMUNITY & LEADERSHIP"]


def test_section_word_in_heading_1_is_still_a_heading(tmp_path):
    doc = Document()
    doc.add_paragraph("Meera Iyer")
    doc.add_heading("Experience", level=1)
    doc.add_paragraph("Analyst, Example Co")
    path = tmp_path / "h1.docx"
    doc.save(str(path))

    outline = read_resume(path)

    assert _by_text(outline, "Experience").is_heading
    assert outline.sections == ["Header", "Experience"]
