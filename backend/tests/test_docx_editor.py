"""
Tests for applying approved changes to a copy of the DOCX (Milestone 5),
and for the validator that checks the result.
"""

import base64
import io

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from app.docx_editor import Change, apply_changes
from app.docx_reader import iter_paragraphs
from app.validator import validate

P6 = "Languages: Python, SQL"
P7 = "Tools: Excel, Power BI, Git"
P10 = "Built a Power BI dashboard used by regional managers to track weekly sales."
P13 = "Inventory Forecast: Forecasted monthly stock using past sales in Python."
P14 = "Portfolio Website: Personal site built with HTML and CSS."

# A 1x1 PNG, for testing lines that contain an image.
TINY_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNk+M9QDwADhgGAWjR9awAAAABJRU5ErkJggg=="
)


def _para(path, pid):
    return {i: p for i, p, _ in iter_paragraphs(Document(str(path)))}[pid]


def _runs(path, pid):
    return [(r.text, r.bold) for r in _para(path, pid).runs if r.text]


def _apply(src, tmp_path, *changes):
    out = tmp_path / "refined.docx"
    result = apply_changes(src, out, list(changes))
    finals = {c.paragraph_id: c.final_text for c in changes
              if c.final_text and c.paragraph_id in {a["paragraph_id"] for a in result.applied}}
    report = validate(src, out, result.applied, finals)
    return out, result, report


# --- Formatting is kept -----------------------------------------------------

def test_bold_label_stays_bold_and_new_words_are_plain(sample_resume, tmp_path):
    out, result, report = _apply(
        sample_resume, tmp_path, Change("s1", "p6", "change", P6, "Languages: SQL, Python")
    )
    assert report.ok, report.to_dict()
    assert _para(out, "p6").text == "Languages: SQL, Python"
    assert _runs(out, "p6") == [("Languages: ", True), ("SQL, Python", None)]


def test_reorder_keeps_bold_label(sample_resume, tmp_path):
    out, _, report = _apply(
        sample_resume, tmp_path, Change("s1", "p7", "change", P7, "Tools: Power BI, Excel, Git")
    )
    assert report.ok
    runs = _runs(out, "p7")
    assert runs[0][1] is True and runs[0][0].strip() == "Tools:"
    assert all(bold is not True for _, bold in runs[1:])


def test_bullet_style_is_kept(sample_resume, tmp_path):
    new = "Built a Power BI dashboard that regional managers used to track weekly sales."
    out, _, report = _apply(sample_resume, tmp_path, Change("s1", "p10", "change", P10, new))
    assert report.ok
    assert _para(out, "p10").style.name == "List Bullet"
    assert _para(out, "p10").text == new


def test_insertion_after_a_word_joins_that_word(sample_resume, tmp_path):
    # " and pandas" continues the previous run (plain), not the bold label.
    new = "Inventory Forecast: Forecasted monthly stock using past sales in Python and pandas."
    out, _, report = _apply(sample_resume, tmp_path, Change("s1", "p13", "change", P13, new))
    assert report.ok
    assert _runs(out, "p13")[0] == ("Inventory Forecast: ", True)
    assert "pandas" in _runs(out, "p13")[1][0]


def test_tabs_are_kept(tmp_path):
    doc = Document()
    p = doc.add_paragraph()
    p.add_run("Example Club | Lead").bold = True
    p.add_run("\t2022 to 2024")
    src = tmp_path / "tabs.docx"
    doc.save(str(src))

    out, _, report = _apply(
        src, tmp_path,
        Change("s1", "p0", "change", "Example Club | Lead\t2022 to 2024", "Example Club | Founding Lead\t2022 to 2024"),
    )
    assert report.ok
    para = _para(out, "p0")
    assert para.text == "Example Club | Founding Lead\t2022 to 2024"
    assert para._p.find(".//" + qn("w:tab")) is not None


def test_hyperlinks_are_kept(tmp_path):
    doc = Document()
    p = doc.add_paragraph("Portfolio: ")
    rel = doc.part.relate_to("https://example.com", "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rel)
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "example.com"
    run.append(text)
    link.append(run)
    p._p.append(link)
    src = tmp_path / "link.docx"
    doc.save(str(src))

    out, _, report = _apply(src, tmp_path, Change("s1", "p0", "change", "Portfolio: example.com", "Website: example.com"))
    assert report.ok
    para = _para(out, "p0")
    assert para.text == "Website: example.com"
    assert para._p.find(qn("w:hyperlink")) is not None


# --- Removing ---------------------------------------------------------------

def test_remove_and_change_together_use_original_ids(sample_resume, tmp_path):
    # Removing p13 must not shift which paragraph "p14" means.
    out, result, report = _apply(
        sample_resume, tmp_path,
        Change("s1", "p13", "remove", P13, None),
        Change("s2", "p14", "change", P14, "Portfolio Website: Site built with HTML and CSS."),
    )
    assert report.ok, report.to_dict()
    texts = [p.text for p in Document(str(out)).paragraphs]
    assert P13 not in texts
    assert "Portfolio Website: Site built with HTML and CSS." in texts


def test_only_line_in_a_table_cell_is_emptied_not_removed(sample_resume, tmp_path):
    out, result, report = _apply(
        sample_resume, tmp_path, Change("s1", "p18", "remove", "Class XII, Example School", None)
    )
    assert report.ok
    assert result.applied[0]["how"] == "emptied"
    assert len(Document(str(out)).tables[0].rows) == 2


# --- Refusals ---------------------------------------------------------------

def test_mismatched_original_is_skipped(sample_resume, tmp_path):
    out, result, report = _apply(
        sample_resume, tmp_path, Change("s1", "p10", "change", "Some other text", "New text")
    )
    assert not result.applied
    assert "doesn't match" in result.skipped[0]["why"]
    assert _para(out, "p10").text == P10


def test_edit_touching_an_image_run_is_skipped(tmp_path):
    doc = Document()
    p = doc.add_paragraph("Phone ")
    p.add_run().add_picture(io.BytesIO(TINY_PNG))
    p.runs[-1]._r.append(OxmlElement("w:t"))
    p.runs[-1]._r[-1].text = "98765"
    src = tmp_path / "img.docx"
    doc.save(str(src))
    original_text = _para(src, "p0").text

    out, result, report = _apply(src, tmp_path, Change("s1", "p0", "change", original_text, "Phone 12345"))
    assert not result.applied
    assert "image" in result.skipped[0]["why"]
    assert report.ok  # nothing changed, and the image is still there


def test_line_with_image_is_not_removed(tmp_path):
    doc = Document()
    p = doc.add_paragraph("Contact ")
    p.add_run().add_picture(io.BytesIO(TINY_PNG))
    src = tmp_path / "img2.docx"
    doc.save(str(src))

    _, result, _ = _apply(src, tmp_path, Change("s1", "p0", "remove", "Contact ", None))
    assert "image" in result.skipped[0]["why"]


def test_section_break_line_is_not_removed(two_column_resume, tmp_path):
    # In the two-column sample, the section break sits on the contact line.
    doc = Document(str(two_column_resume))
    pid, para = next((i, p) for i, p, _ in iter_paragraphs(doc)
                     if p._p.pPr is not None and p._p.pPr.find(qn("w:sectPr")) is not None)
    _, result, _ = _apply(two_column_resume, tmp_path, Change("s1", pid, "remove", para.text, None))
    assert "section" in result.skipped[0]["why"]


# --- The original is never touched -------------------------------------------

def test_original_file_is_unchanged(sample_resume, tmp_path):
    before = sample_resume.read_bytes()
    _apply(sample_resume, tmp_path,
           Change("s1", "p6", "change", P6, "Languages: SQL, Python"),
           Change("s2", "p14", "remove", P14, None))
    assert sample_resume.read_bytes() == before


# --- The validator catches problems -----------------------------------------

def test_validator_catches_an_unapproved_text_change(sample_resume, tmp_path):
    out = tmp_path / "tampered.docx"
    doc = Document(str(sample_resume))
    {i: p for i, p, _ in iter_paragraphs(doc)}["p9"].runs[0].text = "Something else"
    doc.save(str(out))

    report = validate(sample_resume, out, applied=[], final_texts={})
    failed = [c.name for c in report.checks if not c.passed]
    assert "Only approved text changed, nothing went missing" in failed


def test_validator_catches_a_formatting_change(sample_resume, tmp_path):
    out = tmp_path / "tampered.docx"
    doc = Document(str(sample_resume))
    {i: p for i, p, _ in iter_paragraphs(doc)}["p6"].runs[0].bold = False
    doc.save(str(out))

    report = validate(sample_resume, out, applied=[], final_texts={})
    failed = [c.name for c in report.checks if not c.passed]
    assert failed == ["Untouched lines are identical, formatting included"]


def test_validator_catches_a_missing_paragraph(sample_resume, tmp_path):
    out = tmp_path / "tampered.docx"
    doc = Document(str(sample_resume))
    victim = {i: p for i, p, _ in iter_paragraphs(doc)}["p11"]._p
    victim.getparent().remove(victim)
    doc.save(str(out))

    report = validate(sample_resume, out, applied=[], final_texts={})
    assert not report.ok


def test_validator_catches_a_file_that_wont_open(sample_resume, tmp_path):
    out = tmp_path / "broken.docx"
    out.write_bytes(b"not a docx")
    report = validate(sample_resume, out, applied=[], final_texts={})
    assert not report.ok
    assert report.checks[0].name == "The refined file opens"
